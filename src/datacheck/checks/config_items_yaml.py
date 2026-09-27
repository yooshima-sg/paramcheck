"""チェック: 「No」「カテゴリ」「設定項目名」「設定ファイルパス」「設定階層」「パラメータ」の
6項目がすべて列として揃っているシート/テーブルの内容をYAML形式で出力する。

値は次の2つのレイアウトに対応する:
  - Excelテーブル(ListObject)のヘッダー行に必須6項目すべてが揃っている場合:
    そのテーブルの各データ行を抽出する。
  - 通常のシート上に、必須6項目すべてを含むヘッダー行がある場合:
    その直下から、全項目が空になる行までを抽出する。

列名「分類」は「カテゴリ」の別名として扱う（vsftpd設定/Postfix_サービス定義シート等）。

「ビルトイン設定」「EC2パッケージ同梱値」「ROD設定値」「EC2設定値」は任意項目で、
シート/テーブルの絞り込み条件には含めないが、列として存在すれば併せて収集する
（無ければ値はnull）。

出力は1行=1レコードのフラットなリスト。各レコードには元のシート名/テーブル名
（公式テーブルでない場合はnull）を付与してトレーサビリティを確保する。
generate-collect-script で収集したアーカイブ内の実ファイルと突き合わせる際の
基準データとして利用することを想定している。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

NAME = "config-items-yaml"
DESCRIPTION = (
    "「No」「カテゴリ」「設定項目名」「設定ファイルパス」「設定階層」「パラメータ」が"
    "揃ったシート/テーブルの内容を、任意項目（ビルトイン設定/EC2パッケージ同梱値/"
    "ROD設定値/EC2設定値）も含めてYAML形式で出力する"
)

# 絞り込み条件（シート/テーブルが対象になるにはこれらが全部揃っている必要がある）
REQUIRED_COLUMNS = ["No", "カテゴリ", "設定項目名", "設定ファイルパス", "設定階層", "パラメータ"]
# 任意項目（存在すれば収集するが、絞り込み条件には使わない）
OPTIONAL_COLUMNS = ["ビルトイン設定", "EC2パッケージ同梱値", "ROD設定値", "EC2設定値"]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
# 同じ意味で列名が揺れているもの（別名 -> 正式名）。出力キーは正式名に寄せる。
# 例: サーバ証明書/Postfix_サービス定義/vsftpd設定 シートは「カテゴリ」ではなく「分類」。
COLUMN_ALIASES = {"分類": "カテゴリ"}


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _find_header_columns(ws: Worksheet, row: int, min_col: int, max_col: int) -> dict[str, int] | None:
    """指定した行の中でALL_COLUMNSの列位置を探す。REQUIRED_COLUMNSが1つでも欠けていればNone。"""
    mapping: dict[str, int] = {}
    for col in range(min_col, max_col + 1):
        header = _normalize_header(ws.cell(row=row, column=col).value)
        header = COLUMN_ALIASES.get(header, header)
        if header in ALL_COLUMNS and header not in mapping:
            mapping[header] = col
    if set(REQUIRED_COLUMNS).issubset(mapping):
        return mapping
    return None


def _build_record(ws: Worksheet, row: int, col_map: dict[str, int], table_name: str | None) -> dict | None:
    record: dict[str, object] = {"sheet": ws.title, "table": table_name}
    has_value = False
    for label in ALL_COLUMNS:
        col = col_map.get(label)
        value = ws.cell(row=row, column=col).value if col is not None else None
        if not _is_blank(value):
            has_value = True
        record[label] = value
    return record if has_value else None


def _table_covered_cells(ws: Worksheet) -> set[tuple[int, int]]:
    covered: set[tuple[int, int]] = set()
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                covered.add((r, c))
    return covered


def _extract_from_tables(ws: Worksheet) -> list[dict]:
    records: list[dict] = []
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        header_row_count = table.headerRowCount if table.headerRowCount is not None else 1
        if not header_row_count:
            continue

        col_map = _find_header_columns(ws, min_row, min_col, max_col)
        if col_map is None:
            continue

        totals_row_count = table.totalsRowCount or 0
        data_start = min_row + header_row_count
        data_end = max_row - totals_row_count

        table_name = table.displayName or table.name
        for row in range(data_start, data_end + 1):
            record = _build_record(ws, row, col_map, table_name)
            if record is not None:
                records.append(record)
    return records


def _extract_from_plain_sheet(ws: Worksheet, covered: set[tuple[int, int]]) -> list[dict]:
    records: list[dict] = []
    max_row = ws.max_row
    max_col = ws.max_column
    if max_row is None or max_col is None:
        return records

    row = 1
    while row <= max_row:
        if any((row, c) in covered for c in range(1, max_col + 1)):
            row += 1
            continue

        col_map = _find_header_columns(ws, row, 1, max_col)
        if col_map is None:
            row += 1
            continue

        data_row = row + 1
        while data_row <= max_row and not any((data_row, c) in covered for c in col_map.values()):
            record = _build_record(ws, data_row, col_map, None)
            if record is None:
                break
            records.append(record)
            data_row += 1
        row = data_row
    return records


def extract_sheet(ws: Worksheet) -> list[dict]:
    covered = _table_covered_cells(ws)
    return _extract_from_tables(ws) + _extract_from_plain_sheet(ws, covered)


def extract_workbook(path: str | Path) -> list[dict]:
    """全シートを走査し、必須6項目が揃ったシート/テーブルの行レコードをフラットなリストで返す。"""
    wb = load_workbook(filename=path, data_only=True, read_only=False)
    try:
        records: list[dict] = []
        for ws in wb.worksheets:
            records.extend(extract_sheet(ws))
        return records
    finally:
        wb.close()


def add_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("excel_path", type=Path, help="対象のExcelファイル(.xlsx/.xlsm)のパス")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="結果を書き出すYAMLファイルのパス（省略時は標準出力）",
    )


def run(args: argparse.Namespace) -> int:
    if not args.excel_path.exists():
        print(f"エラー: ファイルが見つかりません: {args.excel_path}", file=sys.stderr)
        return 1

    records = extract_workbook(args.excel_path)
    text = yaml.dump(records, allow_unicode=True, sort_keys=False, default_flow_style=False)

    if args.output is not None:
        args.output.write_text(text, encoding="utf-8")
        print(f"抽出結果を書き出しました: {args.output}（{len(records)}件）")
    else:
        print(text, end="")
    return 0
