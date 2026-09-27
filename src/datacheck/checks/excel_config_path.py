"""チェック: Excelブックの各シートから「設定ファイルパス」項目の値を抽出する。

値は次の2つのレイアウトに対応する:
  - Excelテーブル(ListObject)の列見出しが対象ラベルと一致する場合:
    その列の各データ行の値を抽出する。
  - 通常のラベルセルが対象ラベルと一致する場合:
    その直下から空白セルまでの連続したセルを値として抽出する。

1つのセルに空白または改行区切りで複数のパスが記載されている場合は、それぞれを
別々の値として分割して抽出する（同じセル/ラベルを指す複数のMatchになる）。

「ー」のようなダッシュ類のみで構成された値（未設定を表す慣用的なマーカー）は、
明らかにファイルパスではないため抽出結果から除外する。
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

NAME = "excel-config-path"
DESCRIPTION = 'Excelブックの各シートから「設定ファイルパス」項目の値を抽出する'

TARGET_LABEL = "設定ファイルパス"


def _normalize(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = text.rstrip(":：")
    return text.strip()


def _is_target_label(value: object) -> bool:
    return _normalize(value) == TARGET_LABEL


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _split_multi_value(value: object) -> list[str]:
    """1つのセルに空白または改行区切りで複数のパスが入っていることがあるため分割する。"""
    return str(value).split()


# 「未設定」等を表すダッシュ類のみで構成された値は、明らかにファイルパスではないため除外する。
# 例: "ー"(U+30FC), "-", "―", "‐", "−", "－" など。
_PLACEHOLDER_CHARS = set("-‐‑‒–—―−ー－")


def _is_placeholder(token: str) -> bool:
    return len(token) > 0 and all(ch in _PLACEHOLDER_CHARS for ch in token)


def _build_merged_map(ws: Worksheet) -> dict[tuple[int, int], object]:
    merged_value: dict[tuple[int, int], object] = {}
    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_value = ws.cell(row=min_row, column=min_col).value
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                merged_value[(r, c)] = top_value
    return merged_value


def _effective_value(ws: Worksheet, merged_map: dict[tuple[int, int], object], row: int, col: int) -> object:
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    return merged_map.get((row, col))


def _table_covered_cells(ws: Worksheet) -> set[tuple[int, int]]:
    covered: set[tuple[int, int]] = set()
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                covered.add((r, c))
    return covered


@dataclass
class Match:
    source: str  # "table" or "label"
    value: object
    value_cell: str
    label_cell: str | None = None
    table_name: str | None = None
    column: str | None = None

    def to_dict(self) -> dict:
        d: dict = {"source": self.source, "value_cell": self.value_cell, "value": self.value}
        if self.label_cell is not None:
            d["label_cell"] = self.label_cell
        if self.table_name is not None:
            d["table_name"] = self.table_name
        if self.column is not None:
            d["column"] = self.column
        return d


def _extract_from_tables(ws: Worksheet) -> list[Match]:
    results: list[Match] = []
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        header_row_count = table.headerRowCount if table.headerRowCount is not None else 1
        if not header_row_count:
            continue

        header_row = min_row
        target_col = None
        for col in range(min_col, max_col + 1):
            if _is_target_label(ws.cell(row=header_row, column=col).value):
                target_col = col
                break
        if target_col is None:
            continue

        totals_row_count = table.totalsRowCount or 0
        data_start = header_row + header_row_count
        data_end = max_row - totals_row_count

        for row in range(data_start, data_end + 1):
            cell = ws.cell(row=row, column=target_col)
            if _is_blank(cell.value):
                continue
            for token in _split_multi_value(cell.value):
                if _is_placeholder(token):
                    continue
                results.append(
                    Match(
                        source="table",
                        value=token,
                        value_cell=cell.coordinate,
                        table_name=table.displayName or table.name,
                        column=TARGET_LABEL,
                    )
                )
    return results


def _extract_from_labels(
    ws: Worksheet,
    merged_map: dict[tuple[int, int], object],
    covered: set[tuple[int, int]],
) -> list[Match]:
    results: list[Match] = []
    for row in ws.iter_rows():
        for cell in row:
            pos = (cell.row, cell.column)
            if pos in covered:
                continue
            label_value = _effective_value(ws, merged_map, cell.row, cell.column)
            if not _is_target_label(label_value):
                continue

            r = cell.row + 1
            col = cell.column
            while True:
                below_val = _effective_value(ws, merged_map, r, col)
                if _is_blank(below_val):
                    break
                below_cell = ws.cell(row=r, column=col)
                for token in _split_multi_value(below_val):
                    if _is_placeholder(token):
                        continue
                    results.append(
                        Match(
                            source="label",
                            value=token,
                            value_cell=below_cell.coordinate,
                            label_cell=cell.coordinate,
                        )
                    )
                r += 1
    return results


def extract_sheet(ws: Worksheet) -> list[Match]:
    covered = _table_covered_cells(ws)
    merged_map = _build_merged_map(ws)
    return _extract_from_tables(ws) + _extract_from_labels(ws, merged_map, covered)


def extract_workbook(path: str | Path) -> dict[str, list[dict]]:
    """Return {sheet_name: [match_dict, ...]} for every sheet with matches."""
    wb = load_workbook(filename=path, data_only=True, read_only=False)
    try:
        result: dict[str, list[dict]] = {}
        for ws in wb.worksheets:
            matches = extract_sheet(ws)
            if matches:
                result[ws.title] = [m.to_dict() for m in matches]
        return result
    finally:
        wb.close()


def add_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("excel_path", type=Path, help="対象のExcelファイル(.xlsx/.xlsm)のパス")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="結果を書き出すJSONファイルのパス（省略時は標準出力）",
    )


def run(args: argparse.Namespace) -> int:
    if not args.excel_path.exists():
        print(f"エラー: ファイルが見つかりません: {args.excel_path}", file=sys.stderr)
        return 1

    result = extract_workbook(args.excel_path)
    text = json.dumps(result, ensure_ascii=False, indent=2, default=str)

    if args.output is not None:
        args.output.write_text(text, encoding="utf-8")
        print(f"抽出結果を書き出しました: {args.output}")
    else:
        print(text)
    return 0
